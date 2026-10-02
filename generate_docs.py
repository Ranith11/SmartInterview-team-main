import sys
import os
import re

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("python-docx not installed.")
    sys.exit(1)

try:
    from pptx import Presentation
    from pptx.util import Inches as PptxInches, Pt as PptxPt
except ImportError:
    print("python-pptx not installed.")
    sys.exit(1)


def generate_docx():
    doc = Document()
    doc.add_heading('Software Requirement Specification (SRS)', 0)

    try:
        with open('Final_SmartInterview_SRS.md', 'r', encoding='utf-8') as f:
            lines = f.readlines()
    except FileNotFoundError:
        print("Final_SmartInterview_SRS.md not found.")
        return

    for line in lines:
        line = line.strip()
        if not line:
            continue
        
        # Super basic markdown parsing
        if line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        else:
            # clean up strong emphasis
            clean_line = line.replace('**', '')
            doc.add_paragraph(clean_line)

    doc.save('SmartInterview_SRS_Final.docx')
    print("Created SmartInterview_SRS_Final.docx")


def generate_pptx():
    prs = Presentation()
    
    try:
        with open('Final_SmartInterview_PPT.md', 'r', encoding='utf-8') as f:
            content = f.read()
    except FileNotFoundError:
        print("Final_SmartInterview_PPT.md not found.")
        return

    slides_content = content.split('---')
    for slide_text in slides_content:
        lines = slide_text.strip().split('\n')
        if not lines or not lines[0]:
            continue
        
        title = "Slide"
        body = []
        for line in lines:
            line = line.strip()
            if not line:
                continue
            if line.startswith('## '):
                title = line[3:]
            elif line.startswith('# '):
                title = line[2:]
            else:
                body.append(line.replace('**', ''))
        
        slide_layout = prs.slide_layouts[1] # title and content
        slide = prs.slides.add_slide(slide_layout)
        
        if slide.shapes.title:
            slide.shapes.title.text = title
            
        if len(slide.shapes.placeholders) > 1:
            tf = slide.shapes.placeholders[1].text_frame
            tf.text = '\n'.join(body)

    prs.save('SmartInterview_Presentation_Final.pptx')
    print("Created SmartInterview_Presentation_Final.pptx")


if __name__ == "__main__":
    print("Generating documents...")
    generate_docx()
    generate_pptx()
    print("Done.")
