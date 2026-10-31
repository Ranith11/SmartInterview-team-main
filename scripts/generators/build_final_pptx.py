import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DOCS_OUT = os.path.join(REPO_ROOT, "docs", "final_deliverables")
DELIV_OUT = os.path.join(REPO_ROOT, "SmartInterview_Deliverables", "04_Presentation_and_Defense")
os.makedirs(DOCS_OUT, exist_ok=True)
os.makedirs(DELIV_OUT, exist_ok=True)

# Image asset paths
IMG_ARCH = os.path.join(REPO_ROOT, "SmartInterview_Deliverables", "05_System_Diagrams", "Architecture.png")
IMG_WORKFLOW = os.path.join(REPO_ROOT, "SmartInterview_Deliverables", "05_System_Diagrams", "Workflow.png")
IMG_ERD = os.path.join(REPO_ROOT, "SmartInterview_Deliverables", "05_System_Diagrams", "ERD.png")
IMG_DASHBOARD = os.path.join(REPO_ROOT, "screenshots", "04_dashboard.png")
IMG_INTERVIEW = os.path.join(REPO_ROOT, "screenshots", "13_live_interview_session.png")
IMG_RESULTS = os.path.join(REPO_ROOT, "screenshots", "11_results_scorecard.png")
IMG_REMEDIATION = os.path.join(REPO_ROOT, "screenshots", "12_remedial_mastery_mode.png")


def create_presentation():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    NAVY = RGBColor(26, 54, 93)       # #1A365D
    BLUE = RGBColor(37, 99, 235)      # #2563EB
    DARK = RGBColor(30, 41, 59)       # #1E293B
    MUTED = RGBColor(100, 116, 139)   # #64748B
    WHITE = RGBColor(255, 255, 255)

    def add_header(slide, title_text, category="SmartInterview — Academic Viva & Technical Defense"):
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        ctf = cat_box.text_frame
        ctf.word_wrap = True
        cp = ctf.paragraphs[0]
        cp.text = category.upper()
        cp.font.size = Pt(10)
        cp.font.bold = True
        cp.font.color.rgb = BLUE

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        ttf = title_box.text_frame
        ttf.word_wrap = True
        tp = ttf.paragraphs[0]
        tp.text = title_text
        tp.font.size = Pt(26)
        tp.font.bold = True
        tp.font.color.rgb = NAVY

    def add_bullets_box(slide, bullets, left, top, width, height):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        p_first = tf.paragraphs[0]
        p_first.text = ""

        for idx, line in enumerate(bullets):
            p = p_first if idx == 0 else tf.add_paragraph()
            p.text = line.strip()

            if line.startswith("• "):
                p.level = 0
                p.font.size = Pt(15)
                p.font.color.rgb = DARK
                p.space_after = Pt(6)
            elif line.startswith("   - ") or line.startswith("  * ") or line.startswith("     "):
                p.level = 1
                p.font.size = Pt(13)
                p.font.color.rgb = DARK
                p.space_after = Pt(4)
            elif line.endswith(":") or line.startswith("###") or "Benchmark" in line or "Scoring" in line:
                p.level = 0
                p.font.bold = True
                p.font.size = Pt(16)
                p.font.color.rgb = BLUE
                p.space_after = Pt(6)
            else:
                p.level = 0
                p.font.size = Pt(14)
                p.font.color.rgb = DARK
                p.space_after = Pt(5)

    def add_split_slide(title_text, bullets, image_path, img_width=Inches(6.0), img_height=None, img_top=Inches(1.6), img_left=Inches(6.8)):
        slide = prs.slides.add_slide(blank_layout)
        add_header(slide, title_text)
        add_bullets_box(slide, bullets, Inches(0.8), Inches(1.5), Inches(5.8), Inches(5.5))

        if image_path and os.path.exists(image_path):
            if img_height:
                slide.shapes.add_picture(image_path, img_left, img_top, width=img_width, height=img_height)
            else:
                slide.shapes.add_picture(image_path, img_left, img_top, width=img_width)
        return slide

    def add_full_slide(title_text, bullets):
        slide = prs.slides.add_slide(blank_layout)
        add_header(slide, title_text)
        add_bullets_box(slide, bullets, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.5))
        return slide

    # ──────────────────────────────────────────────────────────
    # SLIDE 1: Title Slide
    # ──────────────────────────────────────────────────────────
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "SmartInterview"
    s1.shapes.title.text_frame.paragraphs[0].font.size = Pt(42)
    s1.shapes.title.text_frame.paragraphs[0].font.bold = True
    s1.shapes.title.text_frame.paragraphs[0].font.color.rgb = NAVY

    sub = s1.placeholders[1]
    sub.text = (
        "Personalized Technical Interview Preparation in Higher Education\n"
        "Grounding via Hybrid Graph RAG & Deterministic Bloom's Scaffolding\n\n"
        "Team ID: Team G-1392 | Department of CSE (AI & ML)\n"
        "Keshav Memorial Institute of Technology (KMIT), Hyderabad\n"
        "Faculty Mentor: Dr. TVG Sridevi, Professor, CSE"
    )
    for idx, p in enumerate(sub.text_frame.paragraphs):
        p.font.size = Pt(16 if idx > 1 else 19)
        p.font.color.rgb = DARK

    # ──────────────────────────────────────────────────────────
    # SLIDE 2: Literature Gaps & Problem Statement
    # ──────────────────────────────────────────────────────────
    add_full_slide(
        "Literature Gaps & Problem Statement",
        [
            "The Problem in Higher Education Technical Recruitment:",
            "• Massive student cohorts face high anxiety, lack of articulation practice, and high placement attrition.",
            "• Static coding portals (LeetCode, HackerRank) test binary syntax but cannot assess architectural reasoning or verbal dialogue.",
            "Critical Gaps in Foundational 2026 Base Papers:",
            "• Base Paper 1 (Nagarajan et al., Feb 2026 — IJMRR):",
            "   - Built strictly for enterprise recruiters to screen candidates out.",
            "   - Offers zero student diagnostic feedback, no study roadmap, and volatile session state.",
            "• Base Paper 2 (Wahid et al., March 2026 — IJERT):",
            "   - Silent AI: strictly one-way voice (Whisper STT only, candidate must read text questions).",
            "   - Rigid domain dropdowns without custom Job Description (JD) matching.",
            "• Shared Vulnerabilities in LLM Chatbots:",
            "   - Hallucinations on complex system topics, arbitrary 'vibe grading', and chaotic question jumps."
        ]
    )

    # ──────────────────────────────────────────────────────────
    # SLIDE 3: Proposed Solution
    # ──────────────────────────────────────────────────────────
    add_full_slide(
        "Proposed Solution: SmartInterview Ecosystem",
        [
            "An Adaptive, Grounded Mock Interview & Remediation Platform:",
            "• 3 Distinct Institutional Practice Modes:",
            "   - 1. Job-Specific Mode: Set-theoretic Resume-to-JD alignment (Group A Overlap vs. Group B Gaps).",
            "   - 2. Topic-Based Mode: Targeted deep-dive across 8 core Computer Science domains.",
            "   - 3. Semester Syllabus Mode: Dynamic vector isolation for university course syllabi.",
            "• Hybrid Graph RAG Grounding Engine:",
            "   - Combines Dense Semantic ChromaDB, BM25Okapi Lexical matching, and NetworkX Graph Traversal.",
            "   - Eliminates technical hallucinations with 100% Hit@3 accuracy over 402 curated CS chunks.",
            "• Deterministic Cognitive Scaffolding (Bloom's Taxonomy):",
            "   - Finite state machine advancing through Remember, Understand, Apply, Analyze, Evaluate, Create.",
            "• Objective 5-Factor Answer Evaluation Rubric:",
            "   - Replaces subjective LLM grading with verifiable mathematical scoring and S-BERT cosine similarity.",
            "• Remedial Mastery Mode:",
            "   - Diagnoses weak prerequisite concepts and provides 1-click focused retesting."
        ]
    )

    # ──────────────────────────────────────────────────────────
    # SLIDE 4: Decoupled 4-Tier Distributed System Architecture
    # ──────────────────────────────────────────────────────────
    add_split_slide(
        "Decoupled 4-Tier Distributed Architecture",
        [
            "Zero-Budget, High-Throughput Engineering:",
            "• 1. Presentation Tier (Client):",
            "   - React 18 SPA + Vite 6 + Tailwind CSS.",
            "   - HTML5 Web Audio streaming & visualizers.",
            "• 2. Application Tier (Backend Microservices):",
            "   - FastAPI 0.115 ASGI + Python 3.10+ + Uvicorn.",
            "   - Pydantic v2 validation + HS256 JWT auth.",
            "• 3. Persistence & Vector Storage Tier:",
            "   - MySQL 8.0 (InnoDB): 7 relational tables.",
            "   - ChromaDB 0.6.3: 384-dim local HNSW vector index.",
            "   - NetworkX: In-memory prerequisite knowledge graph.",
            "• 4. Cloud Inference & Neural Voice Tier:",
            "   - Groq Cloud LPU: Llama 3.3 70B & Whisper Large-v3.",
            "   - Microsoft Edge-TTS: In-memory byte streaming (<= 1.2s)."
        ],
        image_path=IMG_ARCH,
        img_width=Inches(5.8),
        img_left=Inches(7.0),
        img_top=Inches(1.8)
    )

    # ──────────────────────────────────────────────────────────
    # SLIDE 5: Core Workflow & Context Ingestion
    # ──────────────────────────────────────────────────────────
    add_split_slide(
        "Core Workflow & Context Ingestion Pipeline",
        [
            "Personalized Technical Staging Loop:",
            "• 1. Document Extraction & Entity Recognition:",
            "   - PyMuPDF parses candidate Resume and target Job Description.",
            "   - Regex filters isolate technical skills and project details.",
            "• 2. Set-Theoretic Skill Gap Analysis:",
            "   - Group A (Role Overlap): Skills present in both Resume & JD.",
            "   - Group B (Role Gaps): Skills required by JD but missing in Resume.",
            "• 3. Adaptive Question Selection:",
            "   - Grounds query context in verified knowledge chunks.",
            "   - Tests Group A for depth and Group B for adaptability.",
            "• 4. Audio Transcription & 5-Signal Evaluation:",
            "   - Speech captured via Whisper, evaluated across 5 metrics.",
            "• 5. Bloom State Machine Progression:",
            "   - Scores >= 80% advance level; scores < 50% trigger remediation."
        ],
        image_path=IMG_WORKFLOW,
        img_width=Inches(5.6),
        img_left=Inches(7.1),
        img_top=Inches(1.6)
    )

    # ──────────────────────────────────────────────────────────
    # SLIDE 6: Hybrid Graph RAG Architecture
    # ──────────────────────────────────────────────────────────
    add_full_slide(
        "Hybrid Graph RAG Engine: Multi-Signal Retrieval",
        [
            "Why Vanilla Vector Search Fails:",
            "• Dense vectors capture broad semantic similarity but miss exact acronyms (e.g. 2PL, ACID, B-Tree).",
            "• Dense search cannot trace prerequisite concept hierarchies (e.g., Locks -> Concurrency -> Transactions).",
            "Tri-Retrieval Architecture & Fusion:",
            "• 1. Dense Semantic Retrieval (ChromaDB + all-MiniLM-L6-v2):",
            "   - 384-dimensional cosine index across 402 curated CS topic chunks in 8 domains.",
            "• 2. Lexical Keyword Retrieval (BM25Okapi):",
            "   - Deterministic exact token matching for algorithms, symbols, and architectural terms.",
            "• 3. Knowledge Graph Traversal (NetworkX In-Memory / Neo4j):",
            "   - Directed ontology of 50 CS concepts, prerequisite edges, and Bloom cognitive depths.",
            "• Reciprocal Rank Fusion (RRF, k=60):",
            "   - Combines rankings via RRF(d) = sum(1 / (k + rank_i(d))) with canonical chunk deduplication.",
            "Empirical Benchmark Results (33 Domain Test Queries):",
            "• Hit@1 = 96.97% | Hit@3 = 100.00% | Mean Reciprocal Rank (MRR) = 0.9798 | Retrieval Latency = 24.6 ms."
        ]
    )

    # ──────────────────────────────────────────────────────────
    # SLIDE 7: Relational Data Schema & Entity Architecture
    # ──────────────────────────────────────────────────────────
    add_split_slide(
        "Relational Data Schema & ACID Storage (ERD)",
        [
            "Robust 7-Table Relational Schema (MySQL 8.0):",
            "• `users`: Candidate identity, bcrypt password hash.",
            "• `resumes`: PyMuPDF parsed skills, work experience, education.",
            "• `job_descriptions`: Target JD entities and domain skills.",
            "• `interview_sessions`:",
            "   - Session difficulty, question count, status.",
            "   - Serialized `AdaptiveState` JSON & current Bloom level.",
            "• `interview_questions`: Question text, Bloom tier, skill, RAG chunk.",
            "• `answers`: Candidate audio transcription and response.",
            "• `answer_evaluations`: 5-factor mathematical rubric metrics.",
            "Database Integrity Guarantees:",
            "• Foreign keys with CASCADE on session deletion.",
            "• Strict ACID compliance with InnoDB engine.",
            "• JSON state serialization prevents mid-session data loss."
        ],
        image_path=IMG_ERD,
        img_width=Inches(2.5),
        img_left=Inches(8.5),
        img_top=Inches(1.5)
    )

    # ──────────────────────────────────────────────────────────
    # SLIDE 8: Deterministic Cognitive Scaffolding (Bloom's Taxonomy)
    # ──────────────────────────────────────────────────────────
    add_full_slide(
        "Deterministic Cognitive Scaffolding (Bloom's Taxonomy)",
        [
            "Strict Pedagogical Progression (Replacing Arbitrary LLM Difficulty):",
            "• 6 Cognitive Levels: Remember (1) -> Understand (2) -> Apply (3) -> Analyze (4) -> Evaluate (5) -> Create (6).",
            "• 3 Difficulty Tiers: Easy, Medium, Hard (enforcing 18 distinct cognitive complexity states).",
            "State Machine Transition Rules:",
            "• Score >= 80% (Mastery Demonstrated):",
            "   - Advances Bloom cognitive level (+1) or advances difficulty tier (Easy -> Medium -> Hard).",
            "   - Ceiling protection: At Create (Level 6) + Hard, scores maintain maximum cognitive depth.",
            "• Score 50% - 79% (Adequate Understanding):",
            "   - Maintains current Bloom level and difficulty tier for concept reinforcement.",
            "• Score < 50% (Foundational Remediation):",
            "   - Regresses Bloom cognitive level (-1) or steps down difficulty to rebuild core principles.",
            "   - Floor protection: At Remember (Level 1) + Easy, floors cleanly without negative underflow.",
            "• Weakness-Biased Skill Selection Algorithm:",
            "   - Automatically samples next questions from lowest-scoring skills to eliminate competency blind spots."
        ]
    )

    # ──────────────────────────────────────────────────────────
    # SLIDE 9: Multi-Signal Answer Evaluation
    # ──────────────────────────────────────────────────────────
    add_full_slide(
        "Multi-Signal Answer Evaluation (Eliminating 'Vibe Grading')",
        [
            "Objective 5-Part Mathematical Scoring Rubric:",
            "• 1. Technical Correctness (30% Weight):",
            "   - LLM logical reasoning and factual accuracy verified against retrieved ground truth.",
            "• 2. Completeness (20% Weight):",
            "   - Coverage of all required problem constraints, trade-offs, and edge cases.",
            "• 3. Relevance (20% Weight):",
            "   - Directness, architectural focus, and absence of conversational filler rhetoric.",
            "• 4. Semantic Cosine Similarity (15% Weight):",
            "   - S-BERT (all-MiniLM-L6-v2) cosine distance between candidate answer and ground-truth chunk.",
            "• 5. Explicit Concept Coverage (15% Weight):",
            "   - Deterministic keyword and entity regex matching for mandatory technical concepts.",
            "Why This Multi-Signal Formula?",
            "• Zero subjectivity: Same answer yields identical, reproducible scores across repeated runs.",
            "• Penalizes confident but factually incorrect responses that fool standard LLM judges."
        ]
    )

    # ──────────────────────────────────────────────────────────
    # SLIDE 10: Live Spoken Interview Room & Neural Voice
    # ──────────────────────────────────────────────────────────
    add_split_slide(
        "Live Spoken Interview Room & Voice Architecture",
        [
            "Real-Time Two-Way Conversational Dialogue:",
            "• Text-to-Speech (Microsoft Edge-TTS):",
            "   - High-fidelity en-US-ChristopherNeural voice profile.",
            "   - Synthesized directly into in-memory io.BytesIO buffers.",
            "   - Audio generation latency <= 1.2 seconds (zero disk overhead).",
            "• Speech-to-Text (Groq Whisper Large-v3):",
            "   - High-accuracy neural speech transcription.",
            "   - Prompt-biased with CS technical terminology.",
            "• Interactive Visual Interface:",
            "   - Current Bloom cognitive level tag (e.g., 'Analyze' Level 4).",
            "   - Real-time speech recognition status & audio controls.",
            "   - Fallback text input for accessibility and quiet environments.",
            "   - Sub-2.5s end-to-end conversational turnaround."
        ],
        image_path=IMG_INTERVIEW,
        img_width=Inches(6.0),
        img_left=Inches(6.8),
        img_top=Inches(1.8)
    )

    # ──────────────────────────────────────────────────────────
    # SLIDE 11: Candidate Evaluation Scorecard & Analytics
    # ──────────────────────────────────────────────────────────
    add_split_slide(
        "Comprehensive Results Scorecard & Analytics",
        [
            "Multi-Dimensional Candidate Feedback:",
            "• Overall Session Score & Performance Tier:",
            "   - Overall percentage with Bloom cognitive depth indicators.",
            "• 5-Axis Radar Breakdown Chart:",
            "   - Correctness, Completeness, Relevance, Semantic, Concept Coverage.",
            "• Question-by-Question Granular Transcripts:",
            "   - Candidate response vs. ground truth reference context.",
            "   - Concrete technical strengths and identified gaps.",
            "• Actionable Recommendations:",
            "   - AI-curated study priorities tagged by urgency.",
            "• Automated PDF Export:",
            "   - 1-click downloadable formal technical evaluation report."
        ],
        image_path=IMG_RESULTS,
        img_width=Inches(6.0),
        img_left=Inches(6.8),
        img_top=Inches(1.8)
    )

    # ──────────────────────────────────────────────────────────
    # SLIDE 12: Remedial Mastery Mode & Targeted Upskilling
    # ──────────────────────────────────────────────────────────
    add_split_slide(
        "Remedial Mastery Mode: Turning Gaps into Strengths",
        [
            "From Evaluation to Accelerated Upskilling:",
            "• Root-Cause Knowledge Graph Traversal:",
            "   - When candidate scores < 50%, engine queries NetworkX graph for missing prerequisite concepts.",
            "• Automated AI Curriculum & Study Guide Synthesis:",
            "   - Generates high-yield study notes, core concepts, and 'where candidates typically lose marks'.",
            "• Isolated Prerequisite Knowledge Chunks:",
            "   - Retrieves exact textbook references for targeted review.",
            "• 1-Click Focused Retest Workflow:",
            "   - Candidates can immediately launch a focused mock interview on their weakest topics.",
            "   - Progress is re-assessed to track score recovery."
        ],
        image_path=IMG_REMEDIATION,
        img_width=Inches(6.0),
        img_left=Inches(6.8),
        img_top=Inches(1.8)
    )

    # ──────────────────────────────────────────────────────────
    # SLIDE 13: System Verification & 355-Test Matrix
    # ──────────────────────────────────────────────────────────
    add_full_slide(
        "System Verification: 355-Test Matrix (100% Passing)",
        [
            "Exhaustive Test Coverage Across All 18 Engineering Phases:",
            "• 355 Total Automated Unit, Integration & Validation Tests (87.90s execution, 100% Pass Rate):",
            "   - Phase 1: Dense Vector RAG Baseline & S-BERT embeddings (25 tests)",
            "   - Phase 2-4: Knowledge Graph Schema, Builder, Storage & Canonical Chunk Resolution (69 tests)",
            "   - Phase 5: BM25 Lexical Keyword Retrieval Engine & Tokenization (34 tests)",
            "   - Phase 6-7: Graph Traversal Retrieval & Hybrid Reciprocal Rank Fusion (48 tests)",
            "   - Phase 8: Reranker Validation & Multi-Signal Scoring (36 tests)",
            "   - Phase 9-10: Pipeline Integration & Bloom's Cognitive Progression (39 tests)",
            "   - Phase 11-13: Security Audit, Evaluation Metrics & Golden Dataset Auditing (48 tests)",
            "   - Phase 14-18: Frozen Baseline, Hybrid Evaluation & Edge Case Safeguards (56 tests)",
            "Key Verified Functional Gates:",
            "• TC-01: Duplicate auth rejected | TC-02: Corrupt uploads safely handled | TC-03: Sub-25ms retrieval",
            "• TC-04: Deterministic Bloom state progression | TC-05: Remedial graph prerequisite traversal",
            "• TC-06: In-memory Edge-TTS <= 1.2s | TC-07: 5-signal rubric convergence | TC-08: Starvation safeguards"
        ]
    )

    # ──────────────────────────────────────────────────────────
    # SLIDE 14: Conclusion & Viva Defense
    # ──────────────────────────────────────────────────────────
    s14 = prs.slides.add_slide(prs.slide_layouts[0])
    s14.shapes.title.text = "Conclusion & Viva Q&A"
    s14.shapes.title.text_frame.paragraphs[0].font.size = Pt(38)
    s14.shapes.title.text_frame.paragraphs[0].font.bold = True
    s14.shapes.title.text_frame.paragraphs[0].font.color.rgb = NAVY

    sub14 = s14.placeholders[1]
    sub14.text = (
        "SmartInterview: Grounded, Adaptive, and Deterministic Technical Interview Preparation.\n\n"
        "• Curated 402-Chunk Hybrid Graph RAG eliminating LLM hallucinations.\n"
        "• Deterministic Bloom's Taxonomy cognitive progression.\n"
        "• Objective 5-Signal mathematical evaluation eliminating 'vibe grading'.\n"
        "• Remedial Mastery Mode for targeted upskilling & focused retests.\n"
        "• Low-latency, two-way neural voice interaction at zero infrastructure cost.\n\n"
        "Thank you! We invite questions and technical discussion from the honorable panel."
    )
    for idx, p in enumerate(sub14.text_frame.paragraphs):
        p.font.size = Pt(15 if idx > 1 else 17)
        p.font.color.rgb = DARK

    # Save outputs to both directories
    p1 = os.path.join(DOCS_OUT, "SmartInterview_Presentation_Final.pptx")
    p2 = os.path.join(DELIV_OUT, "SmartInterview_Project_Presentation.pptx")
    prs.save(p1)
    prs.save(p2)
    print(f"[SUCCESS] Wrote presentation to:\n  1. {p1}\n  2. {p2}")


if __name__ == "__main__":
    create_presentation()
