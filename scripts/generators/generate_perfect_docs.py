import sys
import os
import re
import base64
import requests
import time

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("python-docx not installed.")
    sys.exit(1)

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DIAGRAM_DIR = os.path.join(REPO_ROOT, "docs", "diagrams", "mermaid")
OUT_DIR = os.path.join(REPO_ROOT, "docs", "final_deliverables")
os.makedirs(DIAGRAM_DIR, exist_ok=True)
os.makedirs(OUT_DIR, exist_ok=True)

def encode_mermaid(mermaid_text):
    return base64.b64encode(mermaid_text.encode('utf8')).decode('ascii')

def download_mermaid_image(mermaid_text, filename):
    if os.path.exists(filename) and os.path.getsize(filename) > 0:
        print(f"Using local cached diagram {filename}")
        return True
    b64 = encode_mermaid(mermaid_text)
    url = f'https://mermaid.ink/img/{b64}'
    try:
        r = requests.get(url, timeout=5)
        if r.status_code == 200:
            with open(filename, 'wb') as f:
                f.write(r.content)
            return True
        else:
            print(f"Failed to download image from mermaid.ink: {r.status_code}")
    except Exception as e:
        print(f"Exception downloading image: {e}")
    return False

def generate_docx():
    doc = Document()
    doc.add_heading('Software Requirement Specification (SRS)', 0)

    srs_file = os.path.join(REPO_ROOT, 'Final_SmartInterview_SRS.md')
    try:
        with open(srs_file, 'r', encoding='utf-8') as f:
            content = f.read()
    except FileNotFoundError:
        print(f"{srs_file} not found.")
        return

    # Extract mermaid blocks
    mermaid_pattern = re.compile(r'```mermaid\n(.*?)\n```', re.DOTALL)
    
    parts = mermaid_pattern.split(content)
    img_counter = 0

    for i, part in enumerate(parts):
        if i % 2 == 1:
            # This is a mermaid block
            mermaid_text = part.strip()
            filename = os.path.join(DIAGRAM_DIR, f'mermaid_diagram_{img_counter}.png')
            print(f"Downloading mermaid diagram {img_counter}...")
            if download_mermaid_image(mermaid_text, filename):
                try:
                    doc.add_picture(filename, width=Inches(6.0))
                    # add some space
                    doc.add_paragraph()
                except Exception as e:
                    print(f"Error adding picture {filename}: {e}")
            else:
                doc.add_paragraph("[Failed to render UML diagram as image. Source code below:]", style='Intense Quote')
                doc.add_paragraph(mermaid_text)
            img_counter += 1
        else:
            # Standard markdown
            lines = part.split('\n')
            for line in lines:
                line = line.strip()
                if not line:
                    continue
                
                if line.startswith('### '):
                    doc.add_heading(line[4:].replace('**',''), level=3)
                elif line.startswith('## '):
                    doc.add_heading(line[3:].replace('**',''), level=2)
                elif line.startswith('# '):
                    doc.add_heading(line[2:].replace('**',''), level=1)
                elif line.startswith('- '):
                    doc.add_paragraph(line[2:].replace('**',''), style='List Bullet')
                elif line.startswith('|'):
                    # Basic table support skip (or render as text)
                    doc.add_paragraph(line.replace('**',''))
                else:
                    doc.add_paragraph(line.replace('**', ''))

    output_file = os.path.join(OUT_DIR, 'SmartInterview_SRS_Final.docx')
    try:
        doc.save(output_file)
        print(f"Created {output_file}")
    except PermissionError:
        print(f"PERMISSION DENIED: Please close {output_file} in Microsoft Word before running this script.")
        sys.exit(1)


if __name__ == "__main__":
    print("Generating perfect DOCX with native UML images...")
    generate_docx()
    print("Done.")
