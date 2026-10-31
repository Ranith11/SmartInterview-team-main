from docx import Document

def check_document():
    doc_path = 'SmartInterview_SRS_Final_Template.docx'
    try:
        doc = Document(doc_path)
    except Exception as e:
        print(f"Error opening document: {e}")
        return

    print("=== HEADINGS ===")
    for p in doc.paragraphs:
        if p.style and p.style.name and p.style.name.startswith('Heading'):
            print(f"{p.style.name}: {p.text}")
            
    print("\n=== IMAGES / DIAGRAMS ===")
    image_count = 0
    for rel in doc.part.rels.values():
        if "image" in rel.target_ref:
            image_count += 1
    print(f"Total images found in document: {image_count}")

    print("\n=== TABLES ===")
    for i, table in enumerate(doc.tables):
        print(f"Table {i+1}: {len(table.rows)} rows, {len(table.columns)} columns")
        # Print first row as headers
        if table.rows:
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            print(f"  Headers: {headers}")

if __name__ == "__main__":
    check_document()
