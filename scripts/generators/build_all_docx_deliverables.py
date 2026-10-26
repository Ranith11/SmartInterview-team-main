import os
import sys
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))

def sanitize_xml(s):
    if not s:
        return ""
    # Remove control characters invalid in XML 1.0
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]', '', s)

def convert_md_to_docx(md_path, docx_path, doc_title):
    print(f"Building DOCX: {docx_path} from {os.path.basename(md_path)}...")
    doc = Document()
    
    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    title_p = doc.add_paragraph()
    title_run = title_p.add_run(sanitize_xml(doc_title))
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_table = False
    table_lines = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return
        
        parsed_rows = []
        for tl in table_lines:
            raw_cells = [c.strip() for c in tl.strip().strip('|').split('|')]
            if all(set(c) <= set('-: ') for c in raw_cells):
                continue
            parsed_rows.append(raw_cells)

        if parsed_rows:
            col_count = max(len(r) for r in parsed_rows)
            t = doc.add_table(rows=len(parsed_rows), cols=col_count)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(parsed_rows):
                for ci, cell_val in enumerate(row):
                    if ci < col_count:
                        cell = t.cell(ri, ci)
                        clean_text = sanitize_xml(cell_val.replace('**', '').replace('`', ''))
                        cell.text = clean_text
                        if ri == 0:
                            for p in cell.paragraphs:
                                for r in p.runs:
                                    r.font.bold = True
            doc.add_paragraph()
        table_lines = []
        in_table = False

    for line in lines:
        line = line.strip()
        if not line:
            if in_table:
                flush_table()
            continue

        if line.startswith('|'):
            in_table = True
            table_lines.append(line)
            continue
        elif in_table:
            flush_table()

        if line.startswith('!['):
            m = re.match(r'!\[(.*?)\]\((.*?)\)', line)
            if m:
                cap_text = sanitize_xml(m.group(1))
                img_path = m.group(2).strip()
                if img_path.startswith('file:///'):
                    img_path = img_path[8:]
                img_path = os.path.normpath(img_path)
                if os.path.exists(img_path):
                    try:
                        p_img = doc.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = p_img.add_run()
                        run.add_picture(img_path, width=Inches(5.5))
                        if cap_text:
                            p_cap = doc.add_paragraph()
                            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            r_cap = p_cap.add_run(cap_text)
                            r_cap.font.italic = True
                            r_cap.font.size = Pt(9.5)
                            r_cap.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
                    except Exception as e:
                        print(f"Error adding picture {img_path}: {e}")
            continue

        if line.startswith('### '):
            doc.add_heading(sanitize_xml(line[4:].replace('**', '')), level=3)
        elif line.startswith('## '):
            doc.add_heading(sanitize_xml(line[3:].replace('**', '')), level=2)
        elif line.startswith('# '):
            doc.add_heading(sanitize_xml(line[2:].replace('**', '')), level=1)
        elif line.startswith('- ') or line.startswith('* '):
            doc.add_paragraph(sanitize_xml(line[2:].replace('**', '')), style='List Bullet')
        elif line.startswith('*Fig') or line.startswith('*Figure'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(sanitize_xml(line.replace('*', '')))
            r.font.italic = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        elif line.startswith('---'):
            continue
        else:
            doc.add_paragraph(sanitize_xml(line.replace('**', '')))

    if in_table:
        flush_table()

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)
    print(f"[SUCCESS] Saved: {docx_path} ({os.path.getsize(docx_path)} bytes)")

if __name__ == "__main__":
    # Generate Research Paper DOCX
    md_paper = os.path.join(REPO_ROOT, "SmartInterview_Deliverables", "03_Research_Paper", "SmartInterview_Research_Paper_IEEE.md")
    docx_paper1 = os.path.join(REPO_ROOT, "SmartInterview_Deliverables", "03_Research_Paper", "SmartInterview_Research_Paper.docx")
    docx_paper2 = os.path.join(REPO_ROOT, "SmartInterview_Progressive_Updates", "03_Research_Paper", "SmartInterview_Research_Paper_v2.docx")
    convert_md_to_docx(md_paper, docx_paper1, "SmartInterview: IEEE Technical Research Paper")
    convert_md_to_docx(md_paper, docx_paper2, "SmartInterview: IEEE Technical Research Paper")

    # Generate Quick Revision Cheatsheet DOCX
    md_cheat = os.path.join(REPO_ROOT, "SmartInterview_Deliverables", "01_Documentation_and_Viva", "SmartInterview_Quick_Revision_Cheatsheet.md")
    docx_cheat = os.path.join(REPO_ROOT, "SmartInterview_Deliverables", "01_Documentation_and_Viva", "SmartInterview_Quick_Revision_Cheatsheet.docx")
    convert_md_to_docx(md_cheat, docx_cheat, "SmartInterview: Quick-Revision Cheatsheet & Viva Defense Card")

    # Generate 5-Member Presentation Plan DOCX
    md_plan = os.path.join(REPO_ROOT, "SmartInterview_Deliverables", "04_Presentation_and_Defense", "SmartInterview_5_Member_Presentation_Plan.md")
    docx_plan = os.path.join(REPO_ROOT, "SmartInterview_Deliverables", "04_Presentation_and_Defense", "SmartInterview_5_Member_Presentation_Plan.docx")
    convert_md_to_docx(md_plan, docx_plan, "SmartInterview: 5-Member Presentation & Defense Plan")

    # Generate Master Encyclopedia DOCX
    md_encyclopedia = os.path.join(REPO_ROOT, "SmartInterview_Deliverables", "01_Documentation_and_Viva", "SmartInterview_Master_Encyclopedia_and_Viva_Defense.md")
    docx_encyclopedia = os.path.join(REPO_ROOT, "SmartInterview_Deliverables", "01_Documentation_and_Viva", "SmartInterview_Master_Encyclopedia_and_Viva_Defense.docx")
    convert_md_to_docx(md_encyclopedia, docx_encyclopedia, "SmartInterview: Architectural Encyclopedia & Master Viva Defense Guide")
