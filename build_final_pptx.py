from pptx import Presentation
from pptx.util import Inches, Pt
import sys

def build_pptx():
    # Load template
    template_path = 'PS_Project_Presentation_Template_1.0_01Sep2026(1).pptx'
    try:
        prs = Presentation(template_path)
    except Exception as e:
        print(f"Error opening template: {e}")
        return

    # Delete all existing slides in the template to start fresh
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    for sld in slides:
        xml_slides.remove(sld)

    # Slide 1: Title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "SmartInterview"
    slide.placeholders[1].text = "Personalized Technical Interview Preparation\nTeam G-1392 | CSE (AI & ML)"
    
    # Slide 2: Contents
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Contents"
    tf = slide.placeholders[1].text_frame
    tf.text = "1. Introduction\n2. Requirement\n3. System Design\n4. Development Workflow\n5. Project Outcome\n6. Conclusion"

    # Slide 3: Introduction
    slide = prs.slides.add_slide(prs.slide_layouts[2]) # Two content
    slide.shapes.title.text = "Introduction"
    tf1 = slide.placeholders[1].text_frame
    tf1.text = "The Problem:\n- Disconnect between curricula and industry.\n- Generic AI bots hallucinate facts and lack pedagogical structure."
    tf2 = slide.placeholders[2].text_frame
    tf2.text = "Our Solution (SmartInterview):\n- RAG-Grounded generation.\n- Deterministic Bloom's Taxonomy progression.\n- Multi-signal objective evaluation."

    # Slide 4: Requirement
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Requirement & Evaluation Metrics"
    tf = slide.placeholders[1].text_frame
    tf.text = "A multi-dimensional scoring rubric is required to accurately grade candidates:\n\n- Technical Correctness (30%)\n- Completeness (20%)\n- Relevance (20%)\n- Semantic Similarity via S-BERT (15%)\n- Concept Coverage (15%)"

    # Slide 5: Design
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "System Design & Architecture"
    tf = slide.placeholders[1].text_frame
    tf.text = "- Frontend: React.js SPA (Vite, Tailwind)\n- Backend: FastAPI (Python)\n- Database: MySQL 8.0 (Session state) & ChromaDB (Vector RAG)\n- AI Models: Groq Cloud Llama 3 (LLM), Whisper API (STT)\n- Audio: Microsoft edge-tts (Neural Voice)"

    # Slide 6: Development Workflow (Pipeline)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Development & Workflow"
    tf = slide.placeholders[1].text_frame
    tf.text = "1. Context Ingestion: Resume and Job Description parsed via PyMuPDF.\n2. Skill Mapping: Group A (Overlap) and Group B (Gaps) identified.\n3. RAG Retrieval: Cosine similarity finds top-3 reference chunks in ChromaDB.\n4. Difficulty Scaling: Bloom's state machine adjusts tier based on score (>80% advances).\n5. Voice Interaction: LLM prompt synthesized to speech and answer transcribed."

    # Slide 7: Project Outcome
    slide = prs.slides.add_slide(prs.slide_layouts[2])
    slide.shapes.title.text = "Project Outcome & Testing"
    tf1 = slide.placeholders[1].text_frame
    tf1.text = "Functional Achievements:\n- Fully stateless JWT auth.\n- High-speed semantic resume matching.\n- Natural, sub-second latency conversational interviews."
    tf2 = slide.placeholders[2].text_frame
    tf2.text = "Key Validations:\n- ChromaDB successfully retrieves accurate CS context.\n- Adaptive engine accurately advances and regresses based on performance."

    # Slide 8: Conclusion
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Conclusion"
    tf = slide.placeholders[1].text_frame
    tf.text = "SmartInterview successfully integrates Retrieval-Augmented Generation with pedagogical state machines.\n\nBy leveraging vector databases and 5-factor scoring, it provides a highly realistic, responsive, and unbiased interview preparation ecosystem for engineering students."

    # Slide 9: Thank You
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Thank You"
    slide.placeholders[1].text = "Questions?"

    # Save presentation
    output_path = 'SmartInterview_Presentation_Final.pptx'
    prs.save(output_path)
    print(f"Generated {output_path} successfully using template!")

if __name__ == "__main__":
    build_pptx()
